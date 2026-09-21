"""Build the v2.7.0 Notebook 01-05 processing specification workbook."""

from __future__ import annotations

import re
import zipfile

from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .context import RuntimeContext
from .oox import DOCUMENT_AUTHOR, custom_properties_xml
from .reproducible import (
    build_timestamp,
    canonicalise_office_identifiers,
    save_atomically,
    write_package,
)
from .parameters import (
    PARAMETER_COLUMNS,
    STALE_LEASE_GROUP,
    application_condition,
    build_parameter_rows,
)

FONT_NAME = "Yu Gothic UI"
MONO_FONT = "Consolas"

TITLE_FONT = Font(name=FONT_NAME, size=15, bold=True, color="174C88")
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, color="4B5563")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT_NAME, size=10)
MONO = Font(name=MONO_FONT, size=9)
EDIT_FONT = Font(name=FONT_NAME, size=10, color="0F6CBD", bold=True)

HEADER_FILL = PatternFill("solid", fgColor="174C88")
BAND_FILL = PatternFill("solid", fgColor="EEF3F9")
#: Differential formats colour a cell through the pattern *background*, so the
#: conditional-formatting fills use bgColor rather than the solid fgColor the
#: static styles use.
OK_FILL = PatternFill(bgColor="DFF6E3")
NG_FILL = PatternFill(bgColor="FDE0E0")
OK_FONT = Font(name=FONT_NAME, size=10, bold=True, color="0E6027")
NG_FONT = Font(name=FONT_NAME, size=10, bold=True, color="A4262C")
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")

NOTEBOOK_PURPOSE = {
    "Notebook_01": "配布 CSV の契約検証と、Ontology 用 Delta テーブルの発行（Core）",
    "Notebook_02": "Ontology のセマンティックメタデータを一括登録（Core）",
    "Notebook_03": "完全な Ontology 定義を 1 回の API 操作で作成（Optional）",
    "Notebook_04": "Workshop アイテム一式を preview-first で一括構築（Optional）",
    "Notebook_05": "Bronze/Silver/Gold・データ品質・Direct Lake 拡張（Optional）",
}

NOTEBOOK_SCOPE = {
    "Notebook_01": "Core",
    "Notebook_02": "Core",
    "Notebook_03": "Optional",
    "Notebook_04": "Optional",
    "Notebook_05": "Optional",
}


def _style_header(sheet: Worksheet, row: int, columns: int) -> None:
    for index in range(1, columns + 1):
        cell = sheet.cell(row=row, column=index)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 28


def _style_body(sheet: Worksheet, first_row: int, last_row: int, columns: int) -> None:
    for row in range(first_row, last_row + 1):
        for index in range(1, columns + 1):
            cell = sheet.cell(row=row, column=index)
            if cell.font is None or cell.font.name != MONO_FONT:
                cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = WRAP
        if (row - first_row) % 2 == 1:
            for index in range(1, columns + 1):
                sheet.cell(row=row, column=index).fill = BAND_FILL


def _finish(
    sheet: Worksheet,
    widths: dict[str, float],
    *,
    freeze: str,
    filter_ref: str | None,
    print_title_rows: str = "4:4",
    orientation: str = "landscape",
    fit_to_width: int = 1,
    print_scale: int | None = None,
    print_title_cols: str | None = None,
) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = freeze
    if filter_ref:
        sheet.auto_filter.ref = filter_ref
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = orientation
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    if print_scale is None:
        sheet.page_setup.fitToWidth = fit_to_width
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    else:
        # "Fit to one page wide" shrinks without a floor, so a reference sheet with
        # long descriptive columns can end up printing 10 pt body text at 5 pt.
        # An explicit scale keeps the printed size predictable; the sheet may then
        # need two pages across, so the identifying columns repeat on each one.
        sheet.page_setup.fitToWidth = 0
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = False
        sheet.page_setup.scale = print_scale
    sheet.print_options.horizontalCentered = True
    if print_title_rows:
        sheet.print_title_rows = print_title_rows
    if print_title_cols:
        sheet.print_title_cols = print_title_cols
    _print_footer(sheet)


def _print_footer(sheet: Worksheet) -> None:
    """Sheet name and version on the left, page numbers on the right."""
    sheet.oddFooter.left.text = f"&F — &A（v{_FOOTER_VERSION[0]}）"
    sheet.oddFooter.left.size = 8
    sheet.oddFooter.left.font = FONT_NAME
    sheet.oddFooter.right.text = "Page &P / &N"
    sheet.oddFooter.right.size = 8
    sheet.oddFooter.right.font = FONT_NAME


def _print_identity_header(sheet: Worksheet, left: str, right: str = "") -> None:
    """Carry the sheet's identity in the page header rather than in its rows.

    A wide sheet tiles horizontally, and a body row is cut at the tile boundary
    like every other row: the Parameters_Core subtitle arrived on continuation
    pages truncated, and the 64-character notebook digest never printed whole at
    all because it spills out of column B into cells the tile break slices. A
    page header is drawn per printed page, is not bound to any column, and repeats
    on every tile - so identity that must survive tiling lives here instead.
    """
    sheet.oddHeader.left.text = left
    sheet.oddHeader.left.size = 8
    sheet.oddHeader.left.font = FONT_NAME
    if right:
        sheet.oddHeader.right.text = right
        sheet.oddHeader.right.size = 8
        sheet.oddHeader.right.font = FONT_NAME


#: One-element holder so ``_print_footer`` can read the version without threading
#: the runtime context through every sheet builder.
_FOOTER_VERSION = [""]


def build(context: RuntimeContext, output: Path, label_info: bytes, calendar=None) -> dict[str, int]:
    workbook = Workbook()
    stats: dict[str, int] = {}
    _FOOTER_VERSION[0] = context.version

    overview = workbook.active
    overview.title = "Overview"
    _overview_sheet(overview, context)
    stats["Overview"] = overview.max_row

    for key in sorted(context.notebooks):
        sheet = workbook.create_sheet(key.replace("Notebook_", "NB"))
        _notebook_sheet(sheet, context, key)
        stats[sheet.title] = sheet.max_row

    parameters = workbook.create_sheet("Parameters")
    _parameters_sheet(parameters, context)
    stats["Parameters"] = parameters.max_row

    core_parameters = workbook.create_sheet("Parameters_Core")
    _core_parameters_sheet(core_parameters, context)
    stats["Parameters_Core"] = core_parameters.max_row

    contracts = workbook.create_sheet("Contracts")
    _contracts_sheet(contracts, context, calendar)
    stats["Contracts"] = contracts.max_row

    if calendar is not None:
        daily = workbook.create_sheet("IncrementDaily")
        _daily_sheet(daily, context, calendar)
        stats["IncrementDaily"] = daily.max_row

    save_atomically(lambda staged: workbook.save(staged), output)
    _attach_label(output, label_info, context)
    return stats


def _overview_sheet(sheet: Worksheet, context: RuntimeContext) -> None:
    sheet["A1"] = "Furusato Workshop Notebook 01–05 Processing Specification"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        f"Version {context.version} — 配布 Notebook から生成した実際の Cell インベントリ、"
        f"パラメーター仕様、および横断契約。"
    )
    sheet["A2"].font = SUBTITLE_FONT
    sheet["A4"] = "編集方法"
    sheet["A4"].font = EDIT_FONT
    sheet.merge_cells("B4:I4")
    sheet["B4"] = (
        "手順書の指示どおり、Participant ID と confirmation 系パラメーターだけを変更します。"
        "Cell の追加・削除・並べ替えは行いません。"
    )
    sheet["B4"].font = BODY_FONT
    sheet["B4"].alignment = WRAP
    sheet.row_dimensions[4].height = 30

    headers = ["Notebook", "区分", "目的", "Cells", "Markdown", "Code", "Parameters", "File SHA-256", "Version"]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=6, column=index, value=header)
    _style_header(sheet, 6, len(headers))

    row = 7
    for key in sorted(context.notebooks):
        notebook = context.notebooks[key]
        sheet_name = key.replace("Notebook_", "NB")
        last = 5 + len(notebook.cells)
        sheet.cell(row=row, column=1, value=notebook.name)
        sheet.cell(row=row, column=2, value=NOTEBOOK_SCOPE[key])
        sheet.cell(row=row, column=3, value=NOTEBOOK_PURPOSE[key])
        sheet.cell(row=row, column=4, value=f"=COUNTA('{sheet_name}'!A6:A{last})")
        sheet.cell(row=row, column=5, value=f"=COUNTIF('{sheet_name}'!C6:C{last},\"markdown\")")
        sheet.cell(row=row, column=6, value=f"=COUNTIF('{sheet_name}'!C6:C{last},\"code\")")
        sheet.cell(row=row, column=7, value=f"=COUNTIF('Parameters'!A:A,\"{notebook.name}\")")
        hash_cell = sheet.cell(row=row, column=8, value=notebook.sha256)
        hash_cell.font = MONO
        sheet.cell(row=row, column=9, value=notebook.version)
        row += 1

    _style_body(sheet, 7, row - 1, len(headers))
    for index in range(7, row):
        sheet.cell(row=index, column=8).font = MONO
    _finish(
        sheet,
        {"A": 46, "B": 10, "C": 62, "D": 9, "E": 11, "F": 8, "G": 12, "H": 68, "I": 10},
        freeze="A7",
        filter_ref=f"A6:I{row - 1}",
        print_title_rows="6:6",
        print_title_cols="A:A",
        print_scale=85,
            )


def _notebook_sheet(sheet: Worksheet, context: RuntimeContext, key: str) -> None:
    notebook = context.notebooks[key]
    sheet["A1"] = notebook.name
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = f"{NOTEBOOK_PURPOSE[key]}　|　{notebook.relative_path}"
    sheet["A2"].font = SUBTITLE_FONT
    sheet["A3"] = "File SHA-256"
    sheet["A3"].font = BODY_FONT
    sheet["B3"] = notebook.sha256
    sheet["B3"].font = MONO
    # Rows 1-3 are unmerged text in narrow columns, so on paper they spilled into
    # the neighbouring cells and the tile boundary cut the spill mid-token - the
    # Japanese purpose and the tail of the notebook path ended up painted over
    # each other. Merging each identity row across the used width gives it a box
    # that clips instead of colliding; the authoritative copy is the page header.
    for span in ("A1:H1", "A2:H2", "B3:H3"):
        sheet.merge_cells(span)
    sheet["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 26
    _print_identity_header(
        sheet,
        f"{notebook.name}\n{NOTEBOOK_PURPOSE[key]}",
        f"File SHA-256: {notebook.sha256}",
    )

    headers = [
        "Cell",
        "Notebook index",
        "Type",
        "Purpose / heading",
        "First source line",
        "Source lines",
        "Parameter cell",
        "Cell source SHA-256",
    ]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=5, column=index, value=header)
    _style_header(sheet, 5, len(headers))

    row = 6
    for cell in notebook.cells:
        sheet.cell(row=row, column=1, value=cell.number)
        sheet.cell(row=row, column=2, value=cell.index)
        sheet.cell(row=row, column=3, value=cell.cell_type)
        sheet.cell(row=row, column=4, value=cell.heading)
        source_cell = sheet.cell(row=row, column=5, value=cell.first_line)
        source_cell.font = MONO
        sheet.cell(row=row, column=6, value=cell.source_lines)
        sheet.cell(row=row, column=7, value="Yes" if cell.is_parameter_cell else "No")
        hash_cell = sheet.cell(row=row, column=8, value=cell.source_sha256)
        hash_cell.font = MONO
        row += 1

    _style_body(sheet, 6, row - 1, len(headers))
    for index in range(6, row):
        sheet.cell(row=index, column=5).font = MONO
        sheet.cell(row=index, column=8).font = MONO
    _finish(
        sheet,
        {"A": 14, "B": 15, "C": 12, "D": 54, "E": 62, "F": 13, "G": 15, "H": 68},
        freeze="A6",
        filter_ref=f"A5:H{row - 1}",
        print_title_rows="5:5",
        print_title_cols="A:B",
        print_scale=85,
            )
    # Rows 1-3 stay on screen but are excluded from print. Printed, they are a
    # merged box wider than one page, so a horizontal tile slices it and Excel
    # paints the two halves of one string over each other. The page header now
    # carries the name, the purpose and the full digest on every tile, so nothing
    # is lost by starting the printed range at the blank row above the table.
    sheet.print_area = f"A4:H{row - 1}"


def _parameters_sheet(sheet: Worksheet, context: RuntimeContext) -> None:
    sheet["A1"] = "Notebook 01–05 パラメーター仕様"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        "既定値は配布 Notebook のパラメーターセルから取得しています。"
        "参加者が変更するのは「参加者の操作」欄に指示がある項目だけです。"
    )
    sheet["A2"].font = SUBTITLE_FONT

    headers = ["Notebook", "区分", *PARAMETER_COLUMNS]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=4, column=index, value=header)
    _style_header(sheet, 4, len(headers))

    row = 5
    for key in sorted(context.notebooks):
        notebook = context.notebooks[key]
        for values in build_parameter_rows(context, key):
            sheet.cell(row=row, column=1, value=notebook.name)
            sheet.cell(row=row, column=2, value=NOTEBOOK_SCOPE[key])
            for offset, value in enumerate(values, start=3):
                cell = sheet.cell(row=row, column=offset, value=value)
                if offset in (3, 5):
                    cell.font = MONO
            row += 1

    _style_body(sheet, 5, row - 1, len(headers))
    for index in range(5, row):
        sheet.cell(row=index, column=3).font = MONO
        sheet.cell(row=index, column=5).font = MONO
    _finish(
        sheet,
        {
            "A": 46,
            "B": 10,
            "C": 34,
            "D": 12,
            "E": 34,
            "F": 30,
            "G": 52,
            "H": 44,
            "I": 46,
            "J": 14,
        },
        freeze="C5",
        filter_ref=f"A4:J{row - 1}",
    )
    # Ten wide columns never fit one A4 landscape page at a readable scale, so the
    # sheet prints at a fixed scale instead of shrinking without a floor. Rows 1-3
    # carry the title and the scope note, so the print area starts at row 1. At 85 %
    # the body prints above 7 pt but the sheet now runs past one page across, so the
    # three identity columns repeat on every page rather than leaving page 2 as a
    # wall of unlabelled values.
    sheet.page_setup.fitToWidth = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = False
    sheet.page_setup.scale = 85
    sheet.print_area = f"A1:J{row - 1}"
    sheet.print_title_cols = "A:C"


def _core_parameters_sheet(sheet: Worksheet, context: RuntimeContext) -> None:
    sheet["A1"] = "Core で参加者が変更するパラメーター"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        "Core（Notebook 01 と Notebook 02）で実際に手を入れる項目だけを抜き出したチェックリストです。"
        "「適用条件」が「常に適用」の行は毎回設定します。"
        "stale lease 回復時のみの 3 行は、通常の実行では触りません。"
    )
    sheet["A2"].font = SUBTITLE_FONT
    # A2 is a full sentence in a 46-wide column. Excel only spills text into the
    # neighbouring cells while they stay empty, and the printed page cuts the
    # spill at the page boundary, so the subtitle arrived on paper truncated.
    # Merging it across the used width gives it a box of its own that prints
    # whole and cannot overprint the column beside it.
    sheet.merge_cells("A2:H2")
    sheet["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 30
    # The merged box is wider than one printed page, so a horizontal tile still
    # cuts it. The complete sentence therefore also rides in the page header,
    # which is redrawn on every tile.
    _print_identity_header(sheet, f"{sheet['A1'].value}\n{sheet['A2'].value}")

    headers = [
        "Notebook",
        "Parameter",
        "Default",
        "適用条件",
        "参加者の操作",
        "安全ゲート",
        "実施節",
        "確認",
    ]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=4, column=index, value=header)
    _style_header(sheet, 4, len(headers))

    row = 5
    for key in ("Notebook_01", "Notebook_02"):
        notebook = context.notebooks[key]
        for values in build_parameter_rows(context, key):
            name, _type, default, action, _meaning, _behaviour, gate, chapter = values
            # Everything a participant may ever set belongs here: the always-applied
            # parameters and the three stale-lease recovery parameters, which are
            # only used together and would otherwise be invisible in this checklist.
            if name not in STALE_LEASE_GROUP and (
                action.startswith("変更しない") or action.startswith("通常は変更しない")
            ):
                continue
            sheet.cell(row=row, column=1, value=notebook.name)
            sheet.cell(row=row, column=2, value=name).font = MONO
            sheet.cell(row=row, column=3, value=default).font = MONO
            sheet.cell(row=row, column=4, value=application_condition(name))
            sheet.cell(row=row, column=5, value=action)
            sheet.cell(row=row, column=6, value=gate)
            sheet.cell(row=row, column=7, value=chapter)
            sheet.cell(row=row, column=8, value="")
            row += 1

    _style_body(sheet, 5, row - 1, len(headers))
    for index in range(5, row):
        sheet.cell(row=index, column=2).font = MONO
        sheet.cell(row=index, column=3).font = MONO
    _finish(
        sheet,
        {"A": 46, "B": 34, "C": 24, "D": 30, "E": 40, "F": 54, "G": 14, "H": 10},
        freeze="A5",
        filter_ref=f"A4:H{row - 1}",
        print_title_cols="A:A",
        print_scale=85,
    )
    # Rows 1-2 stay on screen but are excluded from print for the same reason as
    # the notebook sheets: merged across the used width they are wider than one
    # page, so a horizontal tile slices the box and paints both halves of the
    # sentence on top of each other. The page header carries the complete title
    # and subtitle on every tile instead.
    sheet.print_area = f"A3:H{row - 1}"


def _contracts_sheet(sheet: Worksheet, context: RuntimeContext, calendar=None) -> None:
    expected = context.expected
    increment = context.expected_increment
    contract = context.ontology_contract

    sheet["A1"] = "横断契約（Cross-file contracts）"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = "すべて完全一致で検証します。1 つでも一致しない場合は先へ進みません。"
    sheet["A2"].font = SUBTITLE_FONT

    headers = ["Contract", "Expected", "Used by", "Rule"]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=4, column=index, value=header)
    _style_header(sheet, 4, len(headers))

    rows = [
        ("Workshop version", context.version, "全 Notebook とパッケージメタデータ", "Exact"),
        ("Dataset contract", context.dataset_manifest["datasetVersion"], "Notebook 01 / データマニフェスト", "Exact"),
        ("配布 CSV 総数", context.csv_count, "Notebook 01 / Pipeline", "Exact"),
        ("静的入力ファイル", len(context.seed_files), "Notebook 01", "Exact"),
        ("増分入力ファイル", len(context.increment_files), "Pipeline / Eventhouse", "Exact"),
        ("Ontology 出力テーブル", len(expected["outputTableCounts"]), "Notebook 01 / Ontology", "Exact"),
        ("Entity types", contract["entityTypes"], "Notebook 02 / 03 / 04", "Exact"),
        ("Static properties", contract["staticProperties"], "Notebook 02 / 03 / 04", "Exact"),
        ("Time-series properties", contract["timeseriesProperties"], "Notebook 02 / 03 / 04", "Exact"),
        ("Relationship types", contract["relationshipTypes"], "Notebook 02 / 03 / 04", "Exact"),
        ("Data bindings", contract["dataBindings"], "Ontology", "Exact"),
        ("Contextualizations", contract["contextualizations"], "Ontology", "Exact"),
        ("Metadata objects", context.metadata_object_count, "Notebook 02", "Exact"),
        ("Node total", expected["nodeTotal"], "Ontology", "Exact"),
        ("Edge total", expected["edgeTotal"], "Ontology", "Exact"),
        ("静的 Donation 行数", expected["donationRows"], "Notebook 01 / Lakehouse", "Exact"),
        ("静的 Donation 合計金額 (JPY)", expected["totalDonationAmountYen"], "Notebook 01 / Lakehouse", "Exact"),
        ("Increment raw rows", increment["rawRows"], "Pipeline / Eventhouse / Notebook 05", "Exact"),
        ("Increment unique EventID", increment["uniqueEventIds"], "Eventhouse / Notebook 05", "Exact"),
        ("Increment duplicate EventID", increment["duplicateEventIds"], "Eventhouse / Notebook 05", "Exact"),
        ("Increment raw amount (JPY)", increment["rawAmountYen"], "Eventhouse", "Exact"),
        ("Increment dedup amount (JPY)", increment["deduplicatedAmountYen"], "Notebook 05", "Exact"),
        (
            "観測窓 (UTC)",
            f"{increment['observationWindowUtc']['from']} 〜 {increment['observationWindowUtc']['to']}",
            "Eventhouse / Notebook 05",
            "Exact",
        ),
        ("Data Agent ソース数", len(context.agent_sources), "Data Agent", "Exact"),
        ("Lakehouse 例クエリ数", len(context.agent_fewshots), "Data Agent", "Exact"),
        ("Participant ID format", "001–999", "Notebook 01–05", "3 digits"),
    ]
    if calendar is not None:
        rows.extend(
            [
                (
                    "観測 UTC 日数",
                    f"{calendar.utc_day_count} ({calendar.first_utc_day} – {calendar.last_utc_day})",
                    "Eventhouse / Notebook 05",
                    "Exact, no gaps",
                ),
                (
                    "1 UTC 日あたり raw 行数",
                    f"{calendar.min_rows}–{calendar.max_rows}",
                    "Eventhouse",
                    "Range",
                ),
                (
                    "1 UTC 日あたり重複排除後行数",
                    f"{calendar.dedup_min_rows}–{calendar.dedup_max_rows}",
                    "Notebook 05",
                    "Range",
                ),
                (
                    "重複が集中する UTC 日",
                    f"{calendar.duplicate_day}: raw {calendar.duplicate_day_raw_rows} "
                    f"- extra {calendar.duplicate_extra_rows_on_day} = {calendar.duplicate_day_dedup_rows}",
                    "Eventhouse / Notebook 05",
                    "Exact",
                ),
                (
                    "追加重複行数 (dedup が削除)",
                    calendar.duplicate_extra_rows_on_day,
                    "Eventhouse / Notebook 05",
                    "Exact",
                ),
                (
                    "重複グループ該当行数 (100 EventID x 2)",
                    calendar.duplicate_group_rows_on_day,
                    "Eventhouse / Notebook 05",
                    "Exact",
                ),
                (
                    "JST 換算の暦日数",
                    f"{calendar.jst_day_count} ({calendar.first_jst_day} – {calendar.last_jst_day})",
                    "Reporting only",
                    "Expected, not an error",
                ),
                (
                    f"JST {calendar.last_jst_day} へ繰り上がる行",
                    f"{calendar.jst_rollover_rows} (UTC {calendar.jst_rollover_from_utc} – {calendar.jst_rollover_to_utc})",
                    "Reporting only",
                    "Expected, not an error",
                ),
                (
                    "file 003 PublishedAtUtc",
                    context.expected_increment["publishedAtUtc"][2],
                    "Eventhouse",
                    "Expected, not an error",
                ),
            ]
        )

    # Gates added by the final quality pass. Each one is a value a participant can
    # read off the product and compare, so it belongs in the same contract sheet.
    facts = _quality_gate_rows(context)
    rows.extend(facts)
    row = 5
    for contract_name, value, used_by, rule in rows:
        sheet.cell(row=row, column=1, value=contract_name)
        cell = sheet.cell(row=row, column=2, value=value)
        cell.font = MONO
        sheet.cell(row=row, column=3, value=used_by)
        sheet.cell(row=row, column=4, value=rule)
        row += 1

    _style_body(sheet, 5, row - 1, len(headers))
    for index in range(5, row):
        sheet.cell(row=index, column=2).font = MONO

    check_row = row + 1
    sheet.cell(row=check_row, column=1, value="自動チェック（開いたときに再計算されます）").font = EDIT_FONT
    sheet.cell(row=check_row + 1, column=1, value="Metadata objects = Entity + Static + Time-series + Relationship")
    sheet.cell(
        row=check_row + 1,
        column=2,
        value=f"=IF(SUM(B11,B12,B13,B14)=B17,\"OK\",\"MISMATCH\")",
    )
    sheet.cell(row=check_row + 2, column=1, value="Increment raw = unique + duplicate")
    sheet.cell(row=check_row + 2, column=2, value="=IF(B22=B23+B24,\"OK\",\"MISMATCH\")")
    sheet.cell(row=check_row + 3, column=1, value="配布 CSV 総数 = 静的 + 増分")
    sheet.cell(row=check_row + 3, column=2, value="=IF(B7=B8+B9,\"OK\",\"MISMATCH\")")

    # Gates from the quality pass, expressed as formulas over the rows above so a
    # future edit to a single number is caught by the sheet itself.
    lookup = {name: index for index, (name, *_rest) in enumerate(rows, start=5)}
    gold_row = lookup["Notebook 05 gold fact 行数"]
    donation_row = lookup["静的 Donation 行数"] if "静的 Donation 行数" in lookup else None
    sheet.cell(row=check_row + 4, column=1, value="Notebook 05 gold = 静的 Donation + 8 月一意 EventID")
    if donation_row:
        sheet.cell(
            row=check_row + 4,
            column=2,
            value=f'=IF(B{gold_row}=B{donation_row}+B{lookup["Increment unique EventID"]},"OK","MISMATCH")'
            if "Increment unique EventID" in lookup
            else f'=IF(B{gold_row}={context.expected["donationRows"]}+'
            f'{context.workspace_contract["analyticsExtension"]["dataQuality"]["acceptedDistinctEventIds"]},"OK","MISMATCH")',
        )
    else:
        sheet.cell(
            row=check_row + 4,
            column=2,
            value=f'=IF(B{gold_row}={context.expected["donationRows"]}+'
            f'{context.workspace_contract["analyticsExtension"]["dataQuality"]["acceptedDistinctEventIds"]},"OK","MISMATCH")',
        )
    sheet.cell(row=check_row + 5, column=1, value="Data Agent 選択要素 = Lakehouse + KQL + Ontology")
    sheet.cell(
        row=check_row + 5,
        column=2,
        value=f'=IF(SUM(B{lookup["Data Agent: Lakehouse で選択するテーブル数"]},'
        f'B{lookup["Data Agent: KQL で選択する要素数"]},'
        f'B{lookup["Data Agent: Ontology で選択する Entity 数"]})='
        f'{len(context.expected["outputTableCounts"]) + 1 + context.ontology_contract["entityTypes"]},'
        '"OK","MISMATCH")',
    )
    checks = 6
    for offset in range(0, checks):
        for column in (1, 2):
            cell = sheet.cell(row=check_row + offset, column=column)
            if cell.font is None or cell.font.color is None:
                cell.font = BODY_FONT
            cell.alignment = WRAP

    # Numeric gates print with thousands separators so a six- or seven-digit
    # expected value can be compared against the product at a glance.
    for index in range(5, row):
        value = sheet.cell(row=index, column=2).value
        if isinstance(value, int) and not isinstance(value, bool):
            sheet.cell(row=index, column=2).number_format = "#,##0"

    # Green when a cross-file contract reconciles, red when it does not, so a
    # mismatch is visible without reading the formula.
    check_range = f"B{check_row}:B{check_row + checks - 1}"
    sheet.conditional_formatting.add(
        check_range,
        CellIsRule(operator="equal", formula=['"OK"'], font=OK_FONT, fill=OK_FILL),
    )
    sheet.conditional_formatting.add(
        check_range,
        CellIsRule(operator="equal", formula=['"MISMATCH"'], font=NG_FONT, fill=NG_FILL),
    )

    _finish(
        sheet,
        {"A": 44, "B": 30, "C": 30, "D": 46},
        freeze="A5",
        filter_ref=f"A4:D{row - 1}",
    )


def _quality_gate_rows(context: RuntimeContext) -> list[tuple[str, object, str, str]]:
    """Contract rows a participant can verify directly in the product."""
    objects = context.kql_objects
    materialized = [o for o in objects if o.command == "create-or-alter materialized-view"]
    many_to_one = [r for r in context.relationships if r.cardinality.startswith("many-to-one")]
    analytics = context.workspace_contract["analyticsExtension"]["dataQuality"]
    gold_rows = int(context.expected["donationRows"]) + int(analytics["acceptedDistinctEventIds"])
    return [
        (
            "KQL 管理コマンド数",
            len(objects),
            "Eventhouse (第 11 章)",
            "Exact — 射影テーブル・関数・update policy は存在しない",
        ),
        (
            "Data Agent: Lakehouse で選択するテーブル数",
            len(context.expected["outputTableCounts"]),
            "Data Agent (第 16.2 節)",
            "Exact",
        ),
        (
            "Data Agent: KQL で選択する要素数",
            len(materialized),
            "Data Agent (第 16.2 節)",
            f"Exact — {materialized[0].name} のみ。raw DonationEvents は選ばない",
        ),
        (
            "Data Agent: Ontology で選択する Entity 数",
            context.ontology_contract["entityTypes"],
            "Data Agent (第 16.2 節)",
            "Exact",
        ),
        (
            "many-to-one の Relationship 数",
            len(many_to_one),
            "Ontology (第 10.1 節)",
            "Exact — エッジ数 = Origin のインスタンス数",
        ),
        (
            "452025 の 8 月観測数",
            context.dataset_manifest["expectedIncrement"].get("topObservedCount")
            or _observed_leader(context)[0],
            "Eventhouse / Ontology (第 14.1 節)",
            "Exact",
        ),
        (
            "452025 の 8 月観測金額 (JPY)",
            context.dataset_manifest["expectedIncrement"].get("topObservedAmountYen")
            or _observed_leader(context)[1],
            "Eventhouse / Ontology (第 14.1 節)",
            "Exact",
        ),
        (
            "Notebook 05 gold fact 行数",
            gold_rows,
            "Notebook 05 (付録 D.3)",
            f"Exact — 静的 {context.expected['donationRows']} + 8 月一意 "
            f"{analytics['acceptedDistinctEventIds']}、DataSource 列で区別",
        ),
        (
            "公開版スモークセット",
            "T03 / T07 / T08",
            "Data Agent (第 18.1 節)",
            "Fixed — 曖昧語 / モデル境界 / 粒度の拒否",
        ),
    ]


def _observed_leader(context: RuntimeContext) -> tuple[int, int]:
    """Recompute the August observation leader straight from the packaged CSVs."""
    import collections
    import csv as _csv

    counts: collections.Counter[str] = collections.Counter()
    amounts: collections.Counter[str] = collections.Counter()
    data = context.root / "workshop" / f"v{context.version}" / "data" / "increment"
    for entry in context.increment_files:
        with (data / entry["file"]).open(encoding="utf-8", newline="") as handle:
            for record in _csv.DictReader(handle):
                counts[record["MunicipalityID"]] += 1
                amounts[record["MunicipalityID"]] += int(record["DonationAmountYen"])
    leader = max(amounts.items(), key=lambda item: (item[1], item[0]))[0]
    return counts[leader], amounts[leader]


def _daily_sheet(sheet: Worksheet, context: RuntimeContext, calendar) -> None:
    increment = context.expected_increment
    sheet["A1"] = "2026 年 8 月 増分の UTC 日次分布（検証済み）"
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = (
        f"UTC {calendar.first_utc_day} – {calendar.last_utc_day} の {calendar.utc_day_count} 日、欠測なし。"
        f"1 日あたり raw {calendar.min_rows}–{calendar.max_rows} 行、"
        f"重複排除後 {calendar.dedup_min_rows}–{calendar.dedup_max_rows} 行。"
    )
    sheet["A2"].font = SUBTITLE_FONT
    sheet.merge_cells("A2:E2")
    sheet["A3"] = (
        f"JST は UTC + 9 時間のため、UTC {calendar.jst_rollover_from_utc} 以降の "
        f"{calendar.jst_rollover_rows} 行は JST {calendar.last_jst_day} に繰り上がります。"
        f"file 003 の PublishedAtUtc {increment['publishedAtUtc'][2]} も同じ理由で 9 月表示になります。"
        "どちらも想定どおりで、観測窓外エラーではありません。"
    )
    sheet["A3"].font = SUBTITLE_FONT
    sheet["A3"].alignment = WRAP
    sheet.merge_cells("A3:E3")
    sheet["A4"] = (
        f"{calendar.duplicate_day}: raw {calendar.duplicate_day_raw_rows} 行 − 追加重複 "
        f"{calendar.duplicate_extra_rows_on_day} 行（{calendar.duplicate_event_ids_on_day} EventID・該当 "
        f"{calendar.duplicate_group_rows_on_day} 行）= {calendar.duplicate_day_dedup_rows} 行。"
        "重複する EventID は 2 行ずつ存在し、dedup が取り除くのはその片方だけです。"
    )
    sheet["A4"].font = SUBTITLE_FONT
    sheet["A4"].alignment = WRAP
    sheet.merge_cells("A4:E4")
    # Two wrapped lines at the current column widths. Setting an explicit height
    # keeps the note readable without the 150-240 pt bands the first draft used.
    sheet.row_dimensions[3].height = 30
    sheet.row_dimensions[4].height = 30

    headers = ["UTC day", "Raw rows", "Raw amount (JPY)", "Deduplicated rows", "Note"]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=6, column=index, value=header)
    _style_header(sheet, 6, len(headers))

    row = 7
    for entry in calendar.utc_days:
        sheet.cell(row=row, column=1, value=entry["date"])
        sheet.cell(row=row, column=2, value=entry["rows"])
        sheet.cell(row=row, column=3, value=entry["amount"])
        sheet.cell(row=row, column=4, value=entry["dedupRows"])
        sheet.cell(
            row=row,
            column=5,
            value=(
                f"追加重複 {calendar.duplicate_extra_rows_on_day} 行"
                f"（{calendar.duplicate_event_ids_on_day} EventID・該当 "
                f"{calendar.duplicate_group_rows_on_day} 行）"
                if entry["date"] == calendar.duplicate_day
                else ""
            ),
        )
        row += 1

    _style_body(sheet, 7, row - 1, len(headers))
    last = row - 1
    for index in range(7, row):
        sheet.cell(row=index, column=2).number_format = "#,##0"
        sheet.cell(row=index, column=3).number_format = "#,##0"
        sheet.cell(row=index, column=4).number_format = "#,##0"

    # The summary block sits below the table with its label merged across A:B so
    # no label is clipped by the narrow date column, and every computed value
    # lands in C (and D for the second half of a range) with a number format.
    summary_row = row + 1
    sheet.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=5)
    header_cell = sheet.cell(row=summary_row, column=1, value="サマリー（開くたびに再計算されます）")
    header_cell.font = EDIT_FONT
    header_cell.alignment = TOP

    summary = (
        ("合計 raw 行数", f"=SUM(B7:B{last})", None, "#,##0"),
        ("合計 raw 金額（円）", f"=SUM(C7:C{last})", None, "#,##0"),
        ("合計 重複排除後行数", f"=SUM(D7:D{last})", None, "#,##0"),
        ("1 UTC 日あたり raw 行数（最小 / 最大）", f"=MIN(B7:B{last})", f"=MAX(B7:B{last})", "#,##0"),
        ("対象 UTC 日数", f"=COUNTA(A7:A{last})", None, "#,##0"),
        ("dedup が除去した追加重複行数", f"=SUM(B7:B{last})-SUM(D7:D{last})", None, "#,##0"),
    )
    for offset, (label, first_value, second_value, number_format) in enumerate(summary, start=1):
        target = summary_row + offset
        sheet.merge_cells(start_row=target, start_column=1, end_row=target, end_column=2)
        label_cell = sheet.cell(row=target, column=1, value=label)
        label_cell.font = BODY_FONT
        label_cell.alignment = TOP
        value_cell = sheet.cell(row=target, column=3, value=first_value)
        value_cell.font = BODY_FONT
        value_cell.number_format = number_format
        value_cell.alignment = TOP
        if second_value is not None:
            second_cell = sheet.cell(row=target, column=4, value=second_value)
            second_cell.font = BODY_FONT
            second_cell.number_format = number_format
            second_cell.alignment = TOP

    _finish(
        sheet,
        {"A": 18, "B": 14, "C": 22, "D": 22, "E": 46},
        freeze="A7",
        filter_ref=f"A6:E{last}",
        print_title_rows="6:6",
        orientation="portrait",
    )


def _attach_label(path: Path, label_info: bytes, context: RuntimeContext) -> None:
    """Re-attach the Purview label parts that openpyxl does not preserve."""
    reattach_label(path, label_info, context)


def _scrub_workbook_provenance(blob: bytes) -> bytes:
    """Remove the build machine's fingerprints from ``xl/workbook.xml``.

    Excel records the absolute directory it last saved from in ``x15ac:absPath``
    and a random per-save ``documentId`` in ``xr:revisionPtr``. The first leaks the
    operator's home directory into a published deliverable; the second makes two
    otherwise identical builds differ. Neither carries any workbook content.
    """
    text = blob.decode("utf-8")
    text = re.sub(r"<mc:AlternateContent[^>]*>.*?x15ac:absPath.*?</mc:AlternateContent>", "", text, flags=re.DOTALL)
    text = re.sub(r"<x15ac:absPath\b[^>]*/>", "", text)
    text = re.sub(r'\s+documentId="[^"]*"', "", text)
    text = re.sub(r"<xr:revisionPtr\b[^>]*/>", "", text)
    text = re.sub(r"<xr:revisionPtr\b[^>]*>.*?</xr:revisionPtr>", "", text, flags=re.DOTALL)
    return _prune_unused_namespaces(text).encode("utf-8")


#: Namespace prefixes are only meaningful when some element or attribute uses
#: them. Excel declares the whole co-authoring revision family on every part it
#: rewrites, then leaves the declarations behind when the elements are removed.
_NAMESPACE_DECLARATION = re.compile(r'\s+xmlns:([A-Za-z0-9_]+)="[^"]*"')
_IGNORABLE = re.compile(r'\s+mc:Ignorable="([^"]*)"')


def _prune_unused_namespaces(text: str) -> str:
    """Drop ``xmlns:`` declarations and ``mc:Ignorable`` entries with no user left.

    A dangling declaration is legal XML, but a revision namespace that is
    declared and listed as ignorable while carrying no element is exactly the
    kind of inconsistency that makes Excel offer to repair the workbook on open.
    Pruning happens in one direction only: a prefix is removed when nothing in
    the part references it, so every prefix still in use is preserved untouched.
    """
    body = _NAMESPACE_DECLARATION.sub("", _IGNORABLE.sub("", text))
    used = set(re.findall(r"<\s*/?\s*([A-Za-z0-9_]+):", body))
    used |= set(re.findall(r'\s([A-Za-z0-9_]+):[A-Za-z0-9_]+\s*=\s*"', body))

    def trim_ignorable(match: re.Match[str]) -> str:
        keep = [prefix for prefix in match.group(1).split() if prefix in used]
        return f' mc:Ignorable="{" ".join(keep)}"' if keep else ""

    text = _IGNORABLE.sub(trim_ignorable, text)
    if "mc:Ignorable" in text:
        used.add("mc")

    def trim_declaration(match: re.Match[str]) -> str:
        return match.group(0) if match.group(1) in used else ""

    return _NAMESPACE_DECLARATION.sub(trim_declaration, text)


def reattach_label(path: Path, label_info: bytes, context: RuntimeContext) -> None:
    """Write the Purview label parts, relationship and core properties into ``path``.

    Called after every writer that rebuilds the package (openpyxl on build, Excel
    after the formula recalculation), because both drop the classification parts.
    """
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    parts["docMetadata/LabelInfo.xml"] = label_info
    parts["docProps/custom.xml"] = custom_properties_xml()
    parts["xl/workbook.xml"] = _scrub_workbook_provenance(parts["xl/workbook.xml"])
    # Excel stamps the revision family onto every worksheet part it rewrites, so
    # the same scrub runs across the whole package rather than the workbook alone.
    for name in list(parts):
        if name.startswith("xl/worksheets/") and name.endswith(".xml"):
            parts[name] = _scrub_workbook_provenance(parts[name])
        elif name == "xl/styles.xml":
            parts[name] = _scrub_workbook_provenance(parts[name])

    content_types = parts["[Content_Types].xml"].decode("utf-8")
    if "classificationlabels" not in content_types:
        content_types = content_types.replace(
            "</Types>",
            '<Override PartName="/docMetadata/LabelInfo.xml" '
            'ContentType="application/vnd.ms-office.classificationlabels+xml"/></Types>',
        )
    if "custom-properties+xml" not in content_types:
        content_types = content_types.replace(
            "</Types>",
            '<Override PartName="/docProps/custom.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.custom-properties+xml"/></Types>',
        )
    parts["[Content_Types].xml"] = content_types.encode("utf-8")

    package_rels = parts["_rels/.rels"].decode("utf-8")
    additions = []
    if "classificationlabels" not in package_rels:
        additions.append(
            '<Relationship Id="rIdLabelInfo" '
            'Type="http://schemas.microsoft.com/office/2020/02/relationships/classificationlabels" '
            'Target="docMetadata/LabelInfo.xml"/>'
        )
    if "custom-properties" not in package_rels:
        additions.append(
            '<Relationship Id="rIdCustomProps" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" '
            'Target="docProps/custom.xml"/>'
        )
    if additions:
        package_rels = package_rels.replace("</Relationships>", "".join(additions) + "</Relationships>")
    parts["_rels/.rels"] = package_rels.encode("utf-8")

    now = build_timestamp()
    parts["docProps/core.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        f"<dc:title>Furusato Notebook 01-05 Processing Specification v{context.version}</dc:title>"
        "<dc:subject>Microsoft Fabric notebook processing specification</dc:subject>"
        f"<dc:creator>{DOCUMENT_AUTHOR}</dc:creator>"
        "<cp:keywords>Microsoft Fabric, Fabric IQ, Ontology, Notebook, Parameters, Workshop</cp:keywords>"
        f"<cp:lastModifiedBy>{DOCUMENT_AUTHOR}</cp:lastModifiedBy>"
        f'<dcterms:created xsi:type="dcterms:W3CDTF">{now}</dcterms:created>'
        f'<dcterms:modified xsi:type="dcterms:W3CDTF">{now}</dcterms:modified>'
        "</cp:coreProperties>"
    ).encode("utf-8")

    # Last write on the workbook: every member is re-emitted with the fixed build
    # instant so two builds of unchanged content are byte-identical.
    write_package(path, parts)
    canonicalise_office_identifiers(path)
