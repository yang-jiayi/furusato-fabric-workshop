"""Explicit field-free navigation for the participant guide.

This export does not certify Microsoft Word pagination.  Its contents list uses
real internal hyperlinks and contains no cached page numbers or Office fields.
The ordinary field-based release path remains the default.
"""

from __future__ import annotations

import zipfile
from collections import Counter
from pathlib import Path

from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from .reproducible import write_package

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
FIELD_TAGS = (W + "fldChar", W + "instrText", W + "fldSimple")
TEXT_FONT = "Noto Sans CJK JP"
CODE_FONT = "Noto Sans Mono CJK JP"
FONT_ATTRIBUTES = ("ascii", "hAnsi", "eastAsia", "cs")


def apply_fonts(path: Path):
    """Use installed CJK fonts for this explicitly selected portable export.

    A Latin fallback for Yu Gothic UI can render Japanese table text as missing
    glyphs.  Bind all direct text and style/default text to the CJK family, remove
    competing theme font attributes, and retain a CJK monospace face for code.
    The ordinary Microsoft Word release renderer never calls this function.
    """
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    def is_mono(fonts):
        return fonts is not None and any(
            "mono" in value.lower() or "consolas" in value.lower() or "courier" in value.lower()
            for value in fonts.attrib.values()
        )

    def bind(properties, mono=False):
        fonts = properties.find(W + "rFonts")
        mono = mono or is_mono(fonts)
        if fonts is None:
            fonts = etree.Element(W + "rFonts")
            style = properties.find(W + "rStyle")
            properties.insert(1 if style is not None else 0, fonts)
        for attribute in list(fonts.attrib):
            if etree.QName(attribute).localname.lower().endswith("theme"):
                del fonts.attrib[attribute]
        for attribute in FONT_ATTRIBUTES:
            fonts.set(W + attribute, CODE_FONT if mono else TEXT_FONT)
        return mono

    styles = etree.fromstring(parts["word/styles.xml"])
    defaults = styles.find(W + "docDefaults")
    if defaults is None:
        defaults = etree.Element(W + "docDefaults")
        styles.insert(0, defaults)
    run_defaults = defaults.find(W + "rPrDefault")
    if run_defaults is None:
        run_defaults = etree.Element(W + "rPrDefault")
        defaults.insert(0, run_defaults)
    properties = run_defaults.find(W + "rPr")
    if properties is None:
        properties = etree.SubElement(run_defaults, W + "rPr")
    bind(properties)
    mono_styles = set()
    style_nodes = list(styles.iter(W + "style"))
    for style in style_nodes:
        properties = style.find(W + "rPr")
        if properties is None:
            properties = etree.SubElement(style, W + "rPr")
        if bind(properties):
            mono_styles.add(style.get(W + "styleId"))
    # Include inherited code styles, regardless of the stylesheet's ordering.
    for _ in range(len(style_nodes)):
        added = False
        for style in style_nodes:
            base = style.find(W + "basedOn")
            if base is not None and base.get(W + "val") in mono_styles and style.get(W + "styleId") not in mono_styles:
                mono_styles.add(style.get(W + "styleId"))
                bind(style.find(W + "rPr"), True)
                added = True
        if not added:
            break
    parts["word/styles.xml"] = etree.tostring(styles, xml_declaration=True, encoding="UTF-8", standalone=True)
    runs = 0
    code_runs = 0
    for name, data in list(parts.items()):
        if not name.startswith("word/") or not name.endswith(".xml") or name == "word/styles.xml":
            continue
        tree = etree.fromstring(data)
        changed = False
        for run in tree.iter(W + "r"):
            properties = run.find(W + "rPr")
            if properties is None:
                properties = etree.Element(W + "rPr")
                run.insert(0, properties)
            style = properties.find(W + "rStyle")
            mono = style is not None and style.get(W + "val") in mono_styles
            code_runs += bind(properties, mono)
            runs += 1
            changed = True
        # Paragraph mark properties can override the default for subsequent text.
        for properties in tree.iter(W + "rPr"):
            if properties.getparent().tag == W + "pPr":
                bind(properties)
                changed = True
        if changed:
            parts[name] = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
    write_package(path, parts)
    return {"body": TEXT_FONT, "code": CODE_FONT, "runsBound": runs, "codeRunsBound": code_runs,
            "fontInstallationRequiredForMatchingRender": True}


def inspect_fonts(parts):
    """Validate effective CJK text families without claiming a layout render."""
    def bound(properties):
        fonts = properties.find(W + "rFonts") if properties is not None else None
        if fonts is None or any(etree.QName(key).localname.lower().endswith("theme") for key in fonts.attrib):
            return False
        families = {fonts.get(W + attribute) for attribute in FONT_ATTRIBUTES}
        return len(families) == 1 and families <= {TEXT_FONT, CODE_FONT}

    styles = etree.fromstring(parts["word/styles.xml"])
    default = styles.find(W + "docDefaults/" + W + "rPrDefault/" + W + "rPr")
    valid = bound(default) and all(bound(style.find(W + "rPr")) for style in styles.iter(W + "style"))
    for name, data in parts.items():
        if name.startswith("word/") and name.endswith(".xml") and name != "word/styles.xml":
            tree = etree.fromstring(data)
            valid = valid and all(bound(run.find(W + "rPr")) for run in tree.iter(W + "r"))
    return valid, f"{TEXT_FONT}; code: {CODE_FONT}; renderer font installation required"


def add_arguments(parser):
    parser.add_argument(
        "--word-navigation", choices=("fields", "headings"), default="fields",
        help="Default: refreshed Microsoft Word fields. Participant-only headings: "
             "field-free chapter links; Microsoft Word pagination is not certified.",
    )


def require_mode(mode, *, participant_edition=False, require_acceptance=False):
    if mode not in {"fields", "headings"}:
        raise ValueError("Unknown Word navigation mode")
    if mode == "headings" and not participant_edition:
        raise ValueError("--word-navigation headings requires --participant-edition")
    if mode == "headings" and require_acceptance:
        raise ValueError("Field-free participant navigation cannot satisfy the final-use Word acceptance gate")
    return mode


def bookmark_name(index):
    return f"FurusatoChapter{index:03d}"


def add_bookmark(paragraph, index):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(index))
    start.set(qn("w:name"), bookmark_name(index))
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(index))
    paragraph._p.append(start)
    paragraph._p.append(end)


def chapter_contents(builder, entries):
    """Keep the 29 chapter/appendix entries as persistent TOC1 hyperlinks."""
    builder.document.add_paragraph("目次", style="TOC Heading")
    for index, title in entries:
        paragraph = builder.document.add_paragraph(style="toc 1")
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), bookmark_name(index))
        hyperlink.set(qn("w:history"), "1")
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        style = OxmlElement("w:rStyle")
        style.set(qn("w:val"), "Hyperlink")
        properties.append(style)
        run.append(properties)
        text = OxmlElement("w:t")
        text.text = title
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
    section = builder.document.add_section(WD_SECTION.NEW_PAGE)
    section.header.is_linked_to_previous = True
    section.footer.is_linked_to_previous = True
    builder._drop_cover_page_setup(section)
    builder._collapse(builder.document.paragraphs[-1])


def remove_carried_fields(path: Path):
    """Remove whole field expressions *and* their cached results from all stories.

    The maintained shell carries a footer PAGE field.  Merely removing its
    instruction would leave an old, plausible-looking page number behind.
    """
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    removed = 0
    for name, data in list(parts.items()):
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        tree = etree.fromstring(data)
        changed = False
        for paragraph in tree.iter(W + "p"):
            depth = 0
            for child in list(paragraph):
                if child.tag == W + "fldSimple":
                    paragraph.remove(child)
                    removed += 1
                    changed = True
                    continue
                markers = list(child.iter(W + "fldChar"))
                in_field = depth > 0
                for marker in markers:
                    kind = marker.get(W + "fldCharType")
                    if kind == "begin":
                        depth += 1
                        removed += 1
                    elif kind == "end":
                        depth -= 1
                        if depth < 0:
                            raise ValueError("Malformed carried Word field: " + name)
                if in_field or markers or list(child.iter(W + "instrText")):
                    paragraph.remove(child)
                    changed = True
            if depth:
                raise ValueError("Unclosed carried Word field: " + name)
        for node in list(tree.iter(W + "updateFields")):
            node.getparent().remove(node)
            changed = True
        if any(list(tree.iter(tag)) for tag in FIELD_TAGS):
            raise ValueError("Unsupported remaining Office field: " + name)
        if changed:
            parts[name] = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
    write_package(path, parts)
    return removed


def inspect_navigation(parts):
    """Check each persistent contents entry against its actual heading bookmark."""
    tree = etree.fromstring(parts["word/document.xml"])
    headings = {}
    toc = []
    bookmarks = {}
    ends = Counter()
    for paragraph in tree.iter(W + "p"):
        style = paragraph.find(W + "pPr/" + W + "pStyle")
        value = style.get(W + "val") if style is not None else ""
        text = "".join(node.text or "" for node in paragraph.iter(W + "t"))
        for node in paragraph.iter(W + "bookmarkStart"):
            name = node.get(W + "name")
            bookmarks.setdefault(name, []).append(node.get(W + "id"))
            if value == "Heading1":
                headings[name] = text
        ends.update(node.get(W + "id") for node in paragraph.iter(W + "bookmarkEnd"))
        if value in {"TOC1", "TOC2"} and text.strip():
            links = list(paragraph.iter(W + "hyperlink"))
            toc.append((text, links))
    links_valid = len(toc) == 29 and all(
        len(links) == 1 and links[0].get(W + "anchor") in headings
        and headings[links[0].get(W + "anchor")] == text
        for text, links in toc
    )
    targets = [links[0].get(W + "anchor") for _, links in toc if len(links) == 1]
    links_valid = links_valid and len(set(targets)) == 29 and set(targets) == set(headings)
    start_ids = [identity for ids in bookmarks.values() for identity in ids]
    bookmarks_valid = bool(bookmarks) and all(
        name and len(ids) == 1 and ends[ids[0]] == 1 for name, ids in bookmarks.items()
    ) and len(set(start_ids)) == len(start_ids) and set(ends) == set(start_ids)
    fields = []
    for name, data in parts.items():
        if name.startswith("word/") and name.endswith(".xml"):
            xml = etree.fromstring(data)
            if any(list(xml.iter(tag)) for tag in FIELD_TAGS) or list(xml.iter(W + "updateFields")):
                fields.append(name)
    return {
        "word.headingNavigationTargets": (links_valid, f"{len(toc)} chapter links"),
        "word.headingNavigationBookmarks": (bookmarks_valid, f"{len(bookmarks)} bookmarks"),
        "word.noOfficeFields": (not fields, str(fields)),
    }
