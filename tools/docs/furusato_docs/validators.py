"""Reusable validators for the Furusato v2.7.0 Office deliverables."""

from __future__ import annotations

import hashlib
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Sequence
from xml.etree import ElementTree

from .typography import ascii_parentheses, TEXT_TAGS, TEXT_ATTRIBUTES

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
CT_NS = "{http://schemas.openxmlformats.org/package/2006/content-types}"

EXPECTED_LABEL_ID = "{f42aa342-8706-4288-bd11-ebb85995028c}"
EXPECTED_SITE_ID = "{72f988bf-86f1-41af-91ab-2d7cd011db47}"
EXPECTED_LABEL_METHOD = "Privileged"
EXPECTED_CONTENT_BITS = "0"
EXPECTED_AUTHOR = "Furusato Fabric Workshop"

#: The legacy ``MSIP_Label_*`` values that must agree with ``LabelInfo.xml`` when
#: the legacy set is present at all.
EXPECTED_LABEL_PROPERTIES = {
    "Enabled": "true",
    "Method": EXPECTED_LABEL_METHOD,
    "Name": "General",
    "SiteId": EXPECTED_SITE_ID.strip("{}"),
    "ContentBits": EXPECTED_CONTENT_BITS,
}

#: Labels and tenants that reached these packages by accident. Office stamps the
#: signing machine's own default sensitivity label into ``docProps/custom.xml`` on
#: every COM save, so these must be scrubbed after the last refresh rather than
#: merely avoided at build time. Store only lowercase-GUID SHA-256 fingerprints
#: so the validator does not itself disclose the private identifiers it rejects.
FORBIDDEN_LABEL_GUID_HASHES: frozenset[str] = frozenset({
    "aac77dce110bc8d178d2c3f6599125ceed85985d348233c7dcb961a1c1055235",
    "0ec244a26fc7f236ad8bd876844992414df3606123ba7ee1ee9f1b46e53639d5",
    "c3590dc853d25875a02997ba84d55d9e2f6a4c9c4c7cdb682076cca9defd0745",
    "734d766d7b7c20ac5c9d4fca21e8c2b9bea536c50077216c208bda7cc92d3065",
})
_GUID_PATTERN = re.compile(r"[0-9a-f]{8}-(?:[0-9a-f]{4}-){3}[0-9a-f]{12}", re.IGNORECASE)


def forbidden_label_guid_fingerprints(text: str) -> tuple[str, ...]:
    """Return matching fingerprints without exposing superseded identifiers."""
    fingerprints = {
        hashlib.sha256(match.group().lower().encode("ascii")).hexdigest()
        for match in _GUID_PATTERN.finditer(text)
    }
    return tuple(sorted(fingerprints & FORBIDDEN_LABEL_GUID_HASHES))

FORBIDDEN_PATTERNS: tuple[tuple[str, str], ...] = (
    (r"Eventstream", "Eventstream flow was removed from v2.7.0"),
    (r"ESFurusatoRealtime", "Eventstream item name was removed"),
    (r"Custom endpoint", "Eventstream custom endpoint was removed"),
    (r"core-events_5", "the Core-5 fixture was removed"),
    (r"Core\s*5\s*件", "the Core-5 send step was removed"),
    # The 100-case pack is an internal QA asset under tools/qa and must never be
    # referenced, linked or named in any participant-facing artifact.
    (r"Test\s*100", "the internal 100-case pack must not appear"),
    (r"Ontology_Test_100", "the superseded participant 100-case record must not appear"),
    (r"100\s*問", "the internal 100-case pack must not appear"),
    (r"100\s*ケース", "the internal 100-case pack must not appear"),
    (r"100[-\s]?case", "the internal 100-case pack must not appear"),
    (r"ONT-JA", "internal 100-case case IDs must not appear"),
    (r"caseId", "internal 100-case case IDs must not appear"),
    (r"ontology-tests", "internal QA assets must not appear"),
    (r"ontology[-_]prompt[-_]pack", "internal QA assets must not appear"),
    (r"prompt[-\s]pack", "internal QA assets must not appear"),
    (r"answer-only-mcp", "internal QA assets must not appear"),
    (r"semantic-review-results", "internal QA assets must not appear"),
    (r"tools[/\\]qa", "internal QA assets must not appear"),
    (r"releaseThresholds", "internal release thresholds must not appear"),
    (r"coverageTargets", "internal QA coverage targets must not appear"),
    (r"assertionCatalog", "internal QA catalogs must not appear"),
    (r"forbiddenClaimCatalog", "internal QA catalogs must not appear"),
    (r"evidenceCatalog", "internal QA catalogs must not appear"),
    (r"baselineProfile", "internal QA baseline profile must not appear"),
    (r"HighValueDonationRule", "the high-value notification exercise was removed"),
    (r"高額通知", "the high-value notification exercise was removed"),
    (r"高額寄付.{0,6}通知", "the high-value notification exercise was removed"),
    (r"Ontology Rule", "the Ontology Rule exercise was removed from participant materials"),
    (r"Manage rules", "the Ontology Rule exercise was removed from participant materials"),
    (r"Numeric state", "the Ontology Rule exercise was removed from participant materials"),
    (r"57,?100", "the high-value threshold is no longer part of any participant flow"),
    (r"v2\.6\.0", "stale version reference"),
    (r"v2\.4\.0", "stale version reference"),
    (r"workshop/v2\.6\.0", "stale path reference"),
    (r"\bSAS\b", "SAS tokens are out of scope for v2.7.0"),
    (r"Send-DonationEvents", "the Eventstream producer script was removed"),
)


@dataclass
class Finding:
    level: str
    check: str
    message: str


@dataclass
class Report:
    target: str
    findings: list[Finding] = field(default_factory=list)
    stats: dict[str, object] = field(default_factory=dict)

    def ok(self, check: str, message: str = "") -> None:
        self.findings.append(Finding("PASS", check, message))

    def fail(self, check: str, message: str) -> None:
        self.findings.append(Finding("FAIL", check, message))

    def warn(self, check: str, message: str) -> None:
        self.findings.append(Finding("WARN", check, message))

    @property
    def failures(self) -> list[Finding]:
        return [item for item in self.findings if item.level == "FAIL"]

    @property
    def warnings(self) -> list[Finding]:
        return [item for item in self.findings if item.level == "WARN"]

    @property
    def passed(self) -> bool:
        return not self.failures


# ------------------------------------------------------------------ packaging
def check_package(path: Path, report: Report, *, expect_label: bool = True) -> dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        broken = archive.testzip()
        if broken:
            report.fail("zip.crc", f"corrupt entry: {broken}")
        else:
            report.ok("zip.crc", f"{len(archive.namelist())} parts, all CRCs valid")
        parts = {name: archive.read(name) for name in archive.namelist()}

    if any(name.startswith("EncryptedPackage") for name in parts):
        report.fail("package.encryption", "package is encrypted (EncryptedPackage part present)")
    else:
        report.ok("package.encryption", "standard OOXML, not encrypted")

    macros = [name for name in parts if "vbaProject" in name or name.endswith(".bin") and "vba" in name.lower()]
    if macros:
        report.fail("package.macros", f"macro parts present: {macros}")
    else:
        report.ok("package.macros", "no macro parts")

    # Content types
    content_types = ElementTree.fromstring(parts["[Content_Types].xml"])
    defaults = {node.get("Extension", "").lower() for node in content_types.findall(f"{CT_NS}Default")}
    overrides = {node.get("PartName") for node in content_types.findall(f"{CT_NS}Override")}
    missing_types = []
    for name in parts:
        if name == "[Content_Types].xml" or name.endswith("/"):
            continue
        extension = name.rsplit(".", 1)[-1].lower() if "." in name else ""
        if f"/{name}" in overrides or extension in defaults:
            continue
        missing_types.append(name)
    if missing_types:
        report.fail("package.contentTypes", f"parts without a content type: {missing_types}")
    else:
        report.ok("package.contentTypes", "every part has a content type")

    # Relationships
    dangling: list[str] = []
    for name, blob in parts.items():
        if not name.endswith(".rels"):
            continue
        base = name.rsplit("_rels/", 1)[0]
        root = ElementTree.fromstring(blob)
        for relationship in root.findall(f"{REL_NS}Relationship"):
            if relationship.get("TargetMode") == "External":
                continue
            target = relationship.get("Target", "")
            resolved = _resolve(base, target)
            if resolved not in parts:
                dangling.append(f"{name} -> {target}")
    if dangling:
        report.fail("package.relationships", f"unresolvable relationship targets: {dangling}")
    else:
        report.ok("package.relationships", "every internal relationship target exists")

    if expect_label:
        check_classification(parts, report)

    return parts


def check_classification(parts: dict[str, bytes], report: Report) -> None:
    """Assert exactly one Purview classification, stated consistently.

    ``docMetadata/LabelInfo.xml`` is authoritative. The legacy ``MSIP_Label_*``
    custom properties are optional, but if present they must describe the *same*
    label on the *same* tenant: Office stamps the signing machine's default label
    into every COM save, and a package that carries two different label GUIDs is
    reported inconsistently by different Office versions.
    """
    label = parts.get("docMetadata/LabelInfo.xml")
    if label is None:
        report.fail("label.part", "docMetadata/LabelInfo.xml is missing")
        return

    text = label.decode("utf-8")
    ids = re.findall(r'\bid="\{([0-9a-fA-F-]{36})\}"', text)
    sites = re.findall(r'siteId="\{([0-9a-fA-F-]{36})\}"', text)
    methods = re.findall(r'method="(\w+)"', text)
    bits = re.findall(r'contentBits="(\d+)"', text)
    if (
        ids == [EXPECTED_LABEL_ID.strip("{}")]
        and sites == [EXPECTED_SITE_ID.strip("{}")]
        and methods == [EXPECTED_LABEL_METHOD]
        and bits == [EXPECTED_CONTENT_BITS]
        and 'removed="0"' in text
    ):
        report.ok(
            "label.part",
            f"exactly one label: {EXPECTED_LABEL_ID} / site {EXPECTED_SITE_ID} / "
            f"{EXPECTED_LABEL_METHOD} / contentBits {EXPECTED_CONTENT_BITS}",
        )
    else:
        report.fail(
            "label.part",
            f"LabelInfo does not state exactly the approved label: ids={ids} sites={sites} "
            f"methods={methods} contentBits={bits}",
        )

    if int(bits[0]) if bits else 1:
        report.fail("label.protection", "contentBits is non-zero: the package would be protected")
    else:
        report.ok("label.protection", "contentBits 0: the label marks but does not encrypt the package")

    package_rels = parts.get("_rels/.rels", b"").decode("utf-8")
    if "classificationlabels" in package_rels:
        report.ok("label.relationship", "classification label relationship present")
    else:
        report.fail("label.relationship", "classification label relationship missing")

    custom = parts.get("docProps/custom.xml", b"").decode("utf-8")
    label_ids = sorted(set(re.findall(r"MSIP_Label_([0-9a-fA-F-]{36})_", custom)))
    if not label_ids:
        report.ok(
            "properties.custom",
            "no legacy MSIP custom properties: LabelInfo is the only classification (intentional)",
        )
    elif label_ids == [EXPECTED_LABEL_ID.strip("{}")]:
        missing = [
            key
            for key, value in EXPECTED_LABEL_PROPERTIES.items()
            if f"<vt:lpwstr>{value}</vt:lpwstr>"
            not in _property_value(custom, f"MSIP_Label_{label_ids[0]}_{key}")
        ]
        if missing:
            report.fail("properties.custom", f"legacy MSIP properties disagree with LabelInfo: {missing}")
        else:
            report.ok(
                "properties.custom",
                f"legacy MSIP set matches LabelInfo ({', '.join(EXPECTED_LABEL_PROPERTIES)})",
            )
    else:
        report.fail(
            "properties.custom",
            f"custom.xml carries label ids that are not the approved one: {label_ids}",
        )

    stray = sorted(
        {
            guid.lower()
            for guid in re.findall(
                r"[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}",
                # The fmtid is the OOXML custom-properties format identifier that the
                # specification requires on every property; it is not a label GUID.
                re.sub(r'fmtid="\{[^"]*\}"', "", custom),
            )
        }
        - {EXPECTED_LABEL_ID.strip("{}").lower(), EXPECTED_SITE_ID.strip("{}").lower()}
    )
    if stray:
        report.fail("properties.customGuids", f"custom.xml carries foreign label/site/action GUIDs: {stray}")
    else:
        report.ok("properties.customGuids", "custom.xml carries no label GUID other than the approved pair")


def _property_value(custom: str, name: str) -> str:
    match = re.search(rf'name="{re.escape(name)}"\s*>(.*?)</property>', custom, re.DOTALL)
    return match.group(1) if match else ""


def check_reproducible_package(path: Path, parts: dict[str, bytes], report: Report, *, scope: str) -> None:
    """Every archive member and both core dates must carry the fixed build instant.

    A published deliverable should be verifiable by rebuilding it, which is only
    possible if nothing in the package depends on when or where the build ran.
    """
    import zipfile as _zipfile

    from .reproducible import build_timestamp, zip_date_time

    expected = zip_date_time()
    with _zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
    drifted = [info.filename for info in infos if info.date_time != expected]
    if drifted:
        report.fail(
            f"{scope}.zipTimestamps",
            f"{len(drifted)} of {len(infos)} members carry a build-time timestamp: {drifted[:4]}",
        )
    else:
        report.ok(
            f"{scope}.zipTimestamps",
            f"all {len(infos)} archive members carry the fixed build instant {expected}",
        )

    core = parts.get("docProps/core.xml", b"").decode("utf-8")
    stamp = build_timestamp()
    dates = re.findall(r"<dcterms:(?:created|modified)[^>]*>(.*?)</dcterms:", core)
    if dates and set(dates) == {stamp}:
        report.ok(f"{scope}.coreTimestamps", f"created and modified are both the declared instant {stamp}")
    else:
        report.fail(f"{scope}.coreTimestamps", f"expected both core dates to be {stamp}, found {dates}")


def check_no_build_provenance(parts: dict[str, bytes], report: Report, *, scope: str) -> None:
    """No package may leak the build machine's paths or the operator's identity.

    Excel in particular records the absolute save directory in ``x15ac:absPath``,
    which would publish the operator's home directory to every participant.
    """
    leaks: list[str] = []
    for name, blob in parts.items():
        if not name.endswith((".xml", ".rels")):
            continue
        text = blob.decode("utf-8", "ignore")
        if re.search(r"[A-Za-z]:\\\\?Users\\", text):
            leaks.append(f"{name}: local filesystem path")
        if "absPath" in text:
            leaks.append(f"{name}: x15ac:absPath")
    if leaks:
        report.fail(f"{scope}.provenance", f"build machine fingerprints present: {sorted(set(leaks))}")
    else:
        report.ok(f"{scope}.provenance", "no local paths or absPath entries in any XML part")


def check_authorship(parts: dict[str, bytes], report: Report, *, scope: str) -> None:
    """``dc:creator`` and ``cp:lastModifiedBy`` must both name the workshop."""
    core = parts.get("docProps/core.xml", b"").decode("utf-8")
    creator = re.search(r"<dc:creator>(.*?)</dc:creator>", core, re.DOTALL)
    modifier = re.search(r"<cp:lastModifiedBy>(.*?)</cp:lastModifiedBy>", core, re.DOTALL)
    actual = (creator.group(1) if creator else "", modifier.group(1) if modifier else "")
    if actual == (EXPECTED_AUTHOR, EXPECTED_AUTHOR):
        report.ok(f"{scope}.authorship", f"dc:creator and cp:lastModifiedBy are '{EXPECTED_AUTHOR}'")
    else:
        report.fail(
            f"{scope}.authorship",
            f"expected '{EXPECTED_AUTHOR}' for both, found creator={actual[0]!r} lastModifiedBy={actual[1]!r}",
        )


def check_no_foreign_label_guids(parts: dict[str, bytes], report: Report, *, scope: str) -> None:
    """No XML part anywhere in the package may mention a superseded label or tenant."""
    hits: list[str] = []
    for name, blob in parts.items():
        if not name.endswith((".xml", ".rels")):
            continue
        text = blob.decode("utf-8", "ignore")
        hits.extend(f"{name}:sha256:{value}" for value in forbidden_label_guid_fingerprints(text))
    if hits:
        report.fail(f"{scope}.foreignLabels", f"superseded label/tenant GUID fingerprints present: {sorted(set(hits))}")
    else:
        report.ok(
            f"{scope}.foreignLabels",
            f"none of the {len(FORBIDDEN_LABEL_GUID_HASHES)} superseded label/tenant GUIDs appear in any XML part",
        )


def _resolve(base: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    prefix = base
    while target.startswith("../"):
        target = target[3:]
        prefix = prefix.rstrip("/").rsplit("/", 1)[0] + "/" if "/" in prefix.rstrip("/") else ""
    return f"{prefix}{target}" if prefix else target


# ------------------------------------------------------------- word documents
def check_document(path: Path, report: Report, *, version: str, expected_title_fragment: str) -> dict[str, object]:
    parts = check_package(path, report)
    document = parts["word/document.xml"].decode("utf-8")
    tree = ElementTree.fromstring(parts["word/document.xml"])

    # Tracked changes
    revisions = re.findall(r"<w:(ins|del|moveFrom|moveTo)\b", document)
    if revisions:
        report.fail("document.trackedChanges", f"{len(revisions)} revision marks found")
    else:
        report.ok("document.trackedChanges", "no tracked changes")

    # Core properties
    core = parts["docProps/core.xml"].decode("utf-8")
    if expected_title_fragment in core and version in core:
        report.ok("properties.core", f"title carries v{version}")
    else:
        report.fail("properties.core", "title is missing the expected fragment or version")
    stale = [pattern for pattern, _ in FORBIDDEN_PATTERNS if re.search(pattern, core)]
    if stale:
        report.fail("properties.stale", f"stale references in core properties: {stale}")
    else:
        report.ok("properties.stale", "no stale references in core properties")

    custom = parts.get("docProps/custom.xml", b"").decode("utf-8")
    if not custom or "MSIP_Label_" in custom:
        report.ok("properties.customPresence", "classification custom properties are in their intended state")
    else:
        report.fail("properties.customPresence", "docProps/custom.xml exists but carries no label properties")
    check_authorship(parts, report, scope="properties")
    check_no_foreign_label_guids(parts, report, scope="properties")
    check_no_build_provenance(parts, report, scope="properties")
    check_reproducible_package(path, parts, report, scope="properties")

    # Headings, tables, images, captions
    headings: list[tuple[str, str]] = []
    captions: list[str] = []
    for paragraph in tree.iter(f"{W}p"):
        style = paragraph.find(f"{W}pPr/{W}pStyle")
        name = style.get(f"{W}val") if style is not None else ""
        text = "".join(node.text or "" for node in paragraph.iter(f"{W}t"))
        if name and name.startswith("Heading"):
            headings.append((name, text))
        if name == "Caption":
            captions.append(text)

    tables = list(tree.iter(f"{W}tbl"))
    data_tables = [
        table for table in tables if table.find(f"{W}tblPr/{W}tblCaption") is not None
    ]
    layout_tables = len(tables) - len(data_tables)
    missing_table_alt = [
        table.find(f"{W}tblPr/{W}tblCaption").get(f"{W}val")
        for table in data_tables
        if not (table.find(f"{W}tblPr/{W}tblDescription") is not None)
    ]
    if missing_table_alt:
        report.fail("accessibility.tableAltText", f"data tables without a description: {missing_table_alt}")
    else:
        report.ok(
            "accessibility.tableAltText",
            f"all {len(data_tables)} data tables carry a title and description "
            f"({layout_tables} single-cell callout/code layout tables are excluded)",
        )
    doc_prs = list(tree.iter(f"{WP}docPr"))
    missing_alt = [node.get("name") or "(unnamed)" for node in doc_prs if not (node.get("descr") or "").strip()]
    if missing_alt:
        report.fail("accessibility.altText", f"{len(missing_alt)} images without alt text: {missing_alt[:5]}")
    else:
        report.ok("accessibility.altText", f"all {len(doc_prs)} inline images carry alt text")

    figure_captions = [text for text in captions if text.startswith("図 ")]
    table_captions = [text for text in captions if text.startswith("表 ")]
    if len(figure_captions) == len(doc_prs) - _cover_logo_count(doc_prs):
        report.ok("captions.figures", f"{len(figure_captions)} figure captions for {len(doc_prs)} images")
    else:
        report.fail(
            "captions.figures",
            f"{len(figure_captions)} figure captions vs {len(doc_prs)} images",
        )
    if len(table_captions) == len(data_tables):
        report.ok("captions.tables", f"{len(table_captions)} captions for {len(data_tables)} data tables")
    else:
        report.fail(
            "captions.tables",
            f"{len(table_captions)} table captions vs {len(data_tables)} data tables",
        )

    # Media integrity and reuse
    rels = ElementTree.fromstring(parts["word/_rels/document.xml.rels"])
    media_targets: dict[str, str] = {}
    for relationship in rels.findall(f"{REL_NS}Relationship"):
        if relationship.get("Type", "").endswith("/image"):
            media_targets[relationship.get("Id", "")] = relationship.get("Target", "")
    embeds = [node.get(f"{R}embed") for node in tree.iter(f"{A}blip")]
    unknown = [embed for embed in embeds if embed not in media_targets]
    if unknown:
        report.fail("media.references", f"blip references without an image relationship: {unknown}")
    else:
        report.ok("media.references", f"{len(embeds)} image references all resolve")

    orphans = [
        target
        for rid, target in media_targets.items()
        if rid not in embeds and f"word/{target}" in parts and "header" not in target
    ]
    if orphans:
        report.warn("media.orphans", f"{len(orphans)} image parts are not referenced from the body")
    else:
        report.ok("media.orphans", "no orphan image parts")

    duplicates: dict[str, int] = {}
    for embed in embeds:
        target = media_targets.get(embed, embed)
        duplicates[target] = duplicates.get(target, 0) + 1
    reused = {target: count for target, count in duplicates.items() if count > 1}
    if reused:
        report.warn("media.reuse", f"image parts used more than once: {reused}")
    else:
        report.ok("media.reuse", "each image part is used exactly once")

    report.stats.update(
        {
            "headings": len(headings),
            "headingLevels": {
                level: sum(1 for name, _ in headings if name == level)
                for level in sorted({name for name, _ in headings})
            },
            "tables": len(tables),
            "dataTables": len(data_tables),
            "layoutTables": layout_tables,
            "tableCaptions": len(table_captions),
            "images": len(doc_prs),
            "figureCaptions": len(figure_captions),
            "paragraphs": len(list(tree.iter(f"{W}p"))),
        }
    )
    check_forbidden_package(parts, report)
    check_parentheses(parts, report)
    return {"text": _document_text(tree), "headings": headings, "parts": parts}


def _cover_logo_count(doc_prs: Sequence) -> int:
    return sum(1 for node in doc_prs if "Fabric IQ" in (node.get("name") or ""))


def _document_text(tree: ElementTree.Element) -> str:
    """Body text with the zero-width word joiners removed.

    The builders insert U+2060 inside a few Japanese technical terms so Word keeps
    them on one line. It carries no meaning, so every content comparison works on
    the text a reader sees.
    """
    return "\n".join(
        "".join(node.text or "" for node in paragraph.iter(f"{W}t"))
        for paragraph in tree.iter(f"{W}p")
    ).replace("\u2060", "")


def check_forbidden(text: str, report: Report, *, check: str = "content.forbidden") -> None:
    hits = []
    for pattern, reason in FORBIDDEN_PATTERNS:
        for match in re.finditer(pattern, text, re.IGNORECASE):
            start = max(0, match.start() - 40)
            hits.append(f"{match.group(0)!r} ({reason}) … {text[start:match.end() + 40]!r}")
            break
    if hits:
        report.fail(check, "; ".join(hits))
    else:
        report.ok(check, f"none of the {len(FORBIDDEN_PATTERNS)} forbidden patterns appear")


#: Package parts that carry authored content. Styles, themes and font tables are
#: excluded because they are inherited verbatim from the style carrier.
AUTHORED_PART_PATTERNS = (
    re.compile(r"^word/(document|header\d*|footer\d*|footnotes|endnotes|comments)\.xml$"),
    re.compile(r"^xl/(workbook|sharedStrings)\.xml$"),
    re.compile(r"^xl/worksheets/.+\.xml$"),
    re.compile(r"^docProps/.+\.xml$"),
    re.compile(r"^docMetadata/.+\.xml$"),
    re.compile(r".+\.rels$"),
)


def check_forbidden_package(parts: dict[str, bytes], report: Report) -> None:
    """Run the forbidden-pattern scan over the raw authored XML.

    Paragraph text alone cannot see image alt text, table descriptions, hyperlink
    targets, headers, footers or document properties, so the internal QA assets are
    banned at the package level as well.

    Word stamps every paragraph and run with a random ``w14:paraId`` /
    ``w14:textId`` hex value. Those are not content and can contain any digit
    sequence, so they are stripped before the scan - otherwise a generated id such
    as ``057100EF`` reads as the retired high-value threshold.
    """
    identifiers = re.compile(r'\sw14:(?:paraId|textId)="[0-9A-Fa-f]+"')
    blob = "\n".join(
        identifiers.sub("", payload.decode("utf-8", "ignore"))
        for name, payload in sorted(parts.items())
        if any(pattern.match(name) for pattern in AUTHORED_PART_PATTERNS)
    )
    check_forbidden(blob, report, check="package.forbidden")


def check_required(text: str, required: Iterable[tuple[str, str]], report: Report, *, check: str) -> None:
    text = ascii_parentheses(text)
    missing = [f"{label}: {value!r}" for label, value in required if ascii_parentheses(value) not in text]
    if missing:
        report.fail(check, f"missing expected values: {missing}")
    else:
        report.ok(check, f"all {len(list(required))} expected values present")


def check_parentheses(parts: dict[str, bytes], report: Report) -> None:
    remaining = []
    for name, content in parts.items():
        if not name.endswith(".xml"):
            continue
        tree = ElementTree.fromstring(content)
        values = []
        for node in tree.iter():
            if node.tag in TEXT_TAGS:
                values.append(node.text or "")
            values.extend(node.get(attribute, "") for attribute in TEXT_ATTRIBUTES.get(node.tag, ()))
        if any(ascii_parentheses(value) != value for value in values):
            remaining.append(name)
    if remaining:
        report.fail("typography.asciiParentheses", f"fullwidth parentheses remain in editable text: {remaining}")
    else:
        report.ok("typography.asciiParentheses", "editable text and accessibility labels use ASCII parentheses")


# ------------------------------------------------------------------- workbook
def check_workbook(path: Path, report: Report) -> dict[str, object]:
    parts = check_package(path, report)
    check_authorship(parts, report, scope="workbook")
    check_no_foreign_label_guids(parts, report, scope="workbook")
    check_no_build_provenance(parts, report, scope="workbook")
    check_reproducible_package(path, parts, report, scope="workbook")
    from openpyxl import load_workbook

    values = load_workbook(path, data_only=False)
    text_chunks: list[str] = []
    formulas = 0
    for sheet in values.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                if isinstance(cell.value, str):
                    text_chunks.append(cell.value)
                    if cell.value.startswith("="):
                        formulas += 1
                else:
                    text_chunks.append(str(cell.value))
    report.stats["sheets"] = [sheet.title for sheet in values.worksheets]
    report.stats["formulaCells"] = formulas
    if formulas:
        report.ok("workbook.formulas", f"{formulas} formula cells recalc on open")
    else:
        report.fail("workbook.formulas", "no formula cells found")

    calc = parts["xl/workbook.xml"].decode("utf-8")
    if "fullCalcOnLoad" in calc or formulas:
        report.ok("workbook.calcChain", "no stale cached values are shipped (openpyxl writes formulas only)")

    for sheet in values.worksheets:
        if sheet.freeze_panes is None:
            report.fail("workbook.freeze", f"{sheet.title} has no freeze pane")
        if sheet.sheet_view.showGridLines:
            report.warn("workbook.gridlines", f"{sheet.title} still shows gridlines")
    if not report.failures:
        report.ok("workbook.freeze", "every sheet has a freeze pane")

    check_forbidden_package(parts, report)
    return {"text": "\n".join(text_chunks), "workbook": values}
