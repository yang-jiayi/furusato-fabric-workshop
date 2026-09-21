"""Parenthesis width is presentation-only in both document formats."""
from __future__ import annotations

import io
import json
from pathlib import Path
import re
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "tools" / "docs"))
sys.path.insert(0, str(ROOT / "tools" / "html"))

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from furusato_docs.docx_kit import DocumentBuilder
from furusato_docs.typography import ascii_parentheses, normalize_word_parentheses
from furusato_docs.validators import Report, check_parentheses, check_required
from furusato_docs.quality import _section
from furusato_html.model import Document as HtmlDocument, Section, Text
from furusato_html import page
from furusato_html.assets import sanitize_svg
from furusato_html.render import RenderContext, esc, rich


class AsciiParenthesesTests(unittest.TestCase):
    def test_only_requested_codepoints_change_without_adding_spaces(self):
        source = "A\uff08B\uff09 (C) \uff21\uff22 \u300cD\u300d"
        expected = "A(B) (C) \uff21\uff22 \u300cD\u300d"
        self.assertEqual(ascii_parentheses(source), expected)
        self.assertEqual(ascii_parentheses(expected), expected)

    def test_word_save_normalizes_all_stories_and_keeps_formatting_and_links(self):
        builder = DocumentBuilder.__new__(DocumentBuilder)
        builder.document = Document()
        builder.hyperlinks = []
        doc = builder.document
        paragraph = doc.add_paragraph()
        run = paragraph.add_run("Body\uff08value\uff09")
        run.bold = True
        heading = doc.add_heading("Heading\uff08value\uff09", level=1)
        doc.sections[0].header.paragraphs[0].text = "Header\uff08value\uff09"
        doc.sections[0].footer.paragraphs[0].text = "Footer\uff08value\uff09"
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Cell\uff08value\uff09"
        caption = OxmlElement("w:tblCaption")
        caption.set(qn("w:val"), "Table\uff08label\uff09")
        table._tbl.tblPr.append(caption)
        image = io.BytesIO()
        Image.new("RGB", (2, 2), "white").save(image, format="PNG")
        image.seek(0)
        shape = doc.add_picture(image)
        shape._inline.docPr.set("descr", "Image\uff08label\uff09")
        image_part = next(part for part in doc.part.package.parts if part.partname.startswith("/word/media/"))
        image_before = image_part.blob
        url = "https://learn.microsoft.com/path\uff08preserve-target\uff09"
        builder._hyperlink(paragraph, url, None)
        doc.core_properties.title = "Title\uff08value\uff09"
        settings = OxmlElement("w:noLineBreaksAfter")
        settings.set(qn("w:val"), "\uff08")
        doc.settings.element.append(settings)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "guide.docx"
            builder.save(path)
            with zipfile.ZipFile(path) as archive:
                parts = {name: archive.read(name) for name in archive.namelist()}
            report = Report("test")
            check_parentheses(parts, report)
            self.assertTrue(report.passed, report.failures)
            loaded = Document(path)
            self.assertIn("Body(value)", loaded.paragraphs[0].text)
            self.assertTrue(loaded.paragraphs[0].runs[0].bold)
            self.assertEqual(loaded.paragraphs[1].style.name, heading.style.name)
            self.assertEqual(loaded.tables[0].cell(0, 0).text, "Cell(value)")
            self.assertEqual(loaded.sections[0].header.paragraphs[0].text, "Header(value)")
            self.assertEqual(loaded.sections[0].footer.paragraphs[0].text, "Footer(value)")
            self.assertEqual(loaded.core_properties.title, "Title(value)")
            targets = [rel.target_ref for rel in loaded.part.rels.values() if rel.is_external]
            self.assertIn(url, targets)
            self.assertEqual(parts[str(image_part.partname).lstrip("/")], image_before)
            self.assertEqual(loaded.settings.element.find(qn("w:noLineBreaksAfter")).get(qn("w:val")), "\uff08")
        before = doc.element.xml
        normalize_word_parentheses(doc)
        self.assertEqual(doc.element.xml, before)

    def test_validators_reject_remaining_parentheses_without_rewriting_expected_content(self):
        report = Report("test")
        parts = {"word/document.xml": (
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:p><w:r><w:t>Text\uff08value\uff09</w:t></w:r></w:p></w:document>'
        ).encode("utf-8")}
        check_parentheses(parts, report)
        self.assertFalse(report.passed)
        expected = [("criterion", "Text\uff08value\uff09")]
        comparisons = Report("test")
        check_required("Text(value)", expected, comparisons, check="content")
        self.assertTrue(comparisons.passed)
        self.assertEqual(expected, [("criterion", "Text\uff08value\uff09")])
        check_required("Text(different)", expected, comparisons, check="negative")
        self.assertFalse(comparisons.passed)

    def test_html_text_attributes_code_and_adjacent_url(self):
        self.assertEqual(esc('A\uff08"x"\uff09'), "A(&quot;x&quot;)")
        markup = rich("See https://learn.microsoft.com/test\uff08note\uff09 and `X\uff08Y\uff09`")
        self.assertIn('href="https://learn.microsoft.com/test"', markup)
        self.assertIn("(note)", markup)
        self.assertIn("<code>X(Y)</code>", markup)
        self.assertNotRegex(markup, "[\uff08\uff09]")

    def test_section_boundaries_remain_exact_across_display_punctuation(self):
        original = "Start\uff08one\uff09\nBody\nEnd\uff08two\uff09\nUnrelated"
        for actual in (original, ascii_parentheses(original)):
            self.assertEqual(
                _section(actual, "Start\uff08one\uff09", "End\uff08two\uff09"),
                "Start(one)\nBody\n",
            )

    def test_svg_accessibility_text_uses_the_same_parentheses(self):
        raw = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1 1"><rect id="box" width="1" height="1"/></svg>'
        result = sanitize_svg(raw, "example", "Title\uff08value\uff09", "Description\uff08value\uff09")
        self.assertNotRegex(result, "[\uff08\uff09]")
        self.assertIn('id="example-box"', result)
        self.assertIn('aria-labelledby="example-title example-desc"', result)
        self.assertIn("Description(value)", result)

    def test_dynamic_ui_and_section_titles_cannot_restore_fullwidth_parentheses(self):
        ui = {key: {"ja": "Label\uff08value\uff09", "en": "Label(value)"}
              for key in ("app.skip", "nav.label", "toc.lead", "top.label")}
        section = Section("ch-1", 1, "1", Text("Title\uff08value\uff09", "Title(value)"))
        document = HtmlDocument([section], [], [], [])
        ctx = RenderContext(ui, None, "2.7.0", "build", "runtime", None, None, {"js": ""})
        replacements = {name: (lambda *args: "") for name in (
            "_head", "_header", "_cover", "_howto", "_contents_panel", "_footer", "_lightbox",
            "render_section", "render_toc",
        )}
        with patch.multiple(page, **replacements):
            result = page.render_page(document, ctx)
        self.assertNotRegex(result, "[\uff08\uff09]")
        strings = json.loads(re.search(r"window.__FURUSATO_STRINGS__=(.*?);window\.", result).group(1))
        titles = json.loads(re.search(r"window.__FURUSATO_SECTIONS__=(.*?);</script>", result).group(1))
        self.assertEqual(strings["app.skip"]["ja"], "Label(value)")
        self.assertEqual(titles[0]["ja"], "Title(value)")
        self.assertEqual(ui["app.skip"]["ja"], "Label\uff08value\uff09")
        self.assertEqual(section.title.ja, "Title\uff08value\uff09")


if __name__ == "__main__":
    unittest.main()
